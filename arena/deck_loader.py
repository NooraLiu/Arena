"""Load characters and region decks from the Arena 牌堆表 .xlsx.

Reads sheets by name and columns by header (so column order can change).
Skips the grey example row and blank rows.

Decks are organized one sheet per zone: 林区牌堆 / 水区牌堆 / 石区牌堆 /
城区牌堆 / 中心牌堆. Each row is one card design; its `张数` is how many
copies go into that zone's deck. `类型` is 武器 / 食物 / 护甲 / 碎片, and
`数值` means +attack (武器), +HP (食物), damage reduced (护甲; a written
"-2" is read as 2) or fragment count (碎片; defaults to 1).
"""
import random
import re
from typing import Dict, List, Optional, Tuple
import openpyxl
from .models import Zone, Card, CardType, Character, RandomEvent

OUTER = [Zone.N, Zone.E, Zone.S, Zone.W]
ALL_ZONES = [Zone.CENTER] + OUTER

_TOKEN_ZONES = {
    # named zones; ring order 林区(N) -> 水区(E) -> 石区(S) -> 城区(W) -> back to 林区
    "林区": [Zone.N], "林": [Zone.N],
    "水区": [Zone.E], "水": [Zone.E],
    "石区": [Zone.S], "石": [Zone.S],
    "城区": [Zone.W], "城": [Zone.W],
    "中心": [Zone.CENTER], "center": [Zone.CENTER], "central": [Zone.CENTER], "c": [Zone.CENTER],
    "上": [Zone.N], "北": [Zone.N], "n": [Zone.N],
    "下": [Zone.S], "南": [Zone.S], "s": [Zone.S],
    "左": [Zone.W], "西": [Zone.W], "w": [Zone.W],
    "右": [Zone.E], "东": [Zone.E], "e": [Zone.E],
    "外区": list(OUTER), "外": list(OUTER),
    "任意": list(ALL_ZONES), "全部": list(ALL_ZONES), "all": list(ALL_ZONES),
}


def parse_zones(text) -> List[Zone]:
    if text is None or str(text).strip() == "":
        return list(OUTER)
    t = str(text).strip().lower()
    for sep in "/、,，;； ":
        t = t.replace(sep, " ")
    result = []
    for tok in t.split():
        for z in _TOKEN_ZONES.get(tok, []):
            if z not in result:
                result.append(z)
    if not result:
        return list(OUTER)
    return sorted(result, key=lambda z: ALL_ZONES.index(z))


def _header_map(ws) -> Dict[str, int]:
    return {str(ws.cell(row=1, column=c).value).strip(): c
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=1, column=c).value is not None}


def _is_data_row(cells: dict, name_key: str) -> bool:
    name = cells.get(name_key)
    if name is None or str(name).strip() == "":
        return False
    joined = " ".join(str(v) for v in cells.values() if v is not None)
    return "示例" not in joined


def _iter_rows(ws, name_key: str):
    headers = _header_map(ws)
    for r in range(2, ws.max_row + 1):
        cells = {h: ws.cell(row=r, column=c).value for h, c in headers.items()}
        if _is_data_row(cells, name_key):
            yield cells


def _int(value, default=0) -> int:
    """Read a number; tolerates text like "+2Health" or "+2 防御" by taking the first integer."""
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return int(value)
    m = re.search(r"[-+]?\d+", str(value))
    return int(m.group()) if m else default


def _open(path):
    return openpyxl.load_workbook(path, data_only=True)


def load_characters(path: str) -> List[Character]:
    wb = _open(path)
    ws = wb["人物"]
    chars = []
    for row in _iter_rows(ws, "角色名"):
        chars.append(Character(
            name=str(row["角色名"]).strip(),
            hp_max=_int(row.get("HP", row.get("HP档(3/5/7)")), 13),
            base_attack=_int(row.get("基础攻击(2/3/4)"), 3),
        ))
    return chars


def _add_cards(decks: Dict[Zone, List[Card]], zones: List[Zone], count: int,
               base_id: str, ctype: CardType, value: int,
               name: str, description: str, effect: str, synergy: str):
    for z in zones:
        for i in range(count):
            decks[z].append(Card(
                id=f"{base_id}-{z.value}{i}", type=ctype, value=value,
                name=name, description=description, effect=effect, synergy=synergy,
            ))


ZONE_SHEETS = [("林区牌堆", Zone.N), ("水区牌堆", Zone.E), ("石区牌堆", Zone.S),
               ("城区牌堆", Zone.W), ("中心牌堆", Zone.CENTER)]

_CARD_TYPES = [("武器", CardType.WEAPON), ("食物", CardType.FOOD),
               ("护甲", CardType.ARMOR), ("碎片", CardType.AMMO)]


def _card_type(text, sheet: str, name: str) -> CardType:
    t = str(text or "").strip()
    for key, ctype in _CARD_TYPES:
        if key in t:
            return ctype
    raise ValueError(f"『{sheet}』里的「{name}」类型是「{t}」,只能填 武器 / 食物 / 护甲 / 碎片。")


def load_decks(path: str, rng: Optional[random.Random] = None) -> Dict[Zone, List[Card]]:
    wb = _open(path)
    decks: Dict[Zone, List[Card]] = {z: [] for z in ALL_ZONES}
    for sheet, zone in ZONE_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        for row in _iter_rows(wb[sheet], "牌名"):
            name = str(row["牌名"]).strip()
            count = _int(row.get("张数"), 1)      # blank = 1; 0 = retired card, kept in the sheet only
            if count <= 0:
                continue
            ctype = _card_type(row.get("类型"), sheet, name)
            value = _int(row.get("数值"), 1 if ctype == CardType.AMMO else 0)
            if ctype == CardType.ARMOR:
                value = abs(value)
            _add_cards(
                decks, [zone], count,
                base_id=str(row.get("ID") or name), ctype=ctype, value=value, name=name,
                description=str(row.get("描述") or ""), effect=str(row.get("特殊效果") or ""),
                synergy=str(row.get("角色协同") or ""),
            )
    if rng is not None:
        for z in decks:
            rng.shuffle(decks[z])
    return decks


def load_events(path: str, rng: Optional[random.Random] = None) -> List[RandomEvent]:
    """Read the 随机事件 sheet into a shuffled event deck (one entry per copy).

    Skips the example row (ID starting with 示例) and the 张数合计 summary row.
    """
    wb = _open(path)
    if "随机事件" not in wb.sheetnames:
        return []
    events: List[RandomEvent] = []
    for row in _iter_rows(wb["随机事件"], "事件名"):
        eid = str(row.get("ID") or "").strip()
        name = str(row.get("事件名") or "").strip()
        if not eid or eid.startswith("示例") or "示例" in name:
            continue
        count = max(1, _int(row.get("张数"), 1))
        for _ in range(count):
            events.append(RandomEvent(
                id=eid, name=name,
                description=str(row.get("描述") or ""),
                target=str(row.get("影响区域") or ""),
                effect=str(row.get("效果") or ""),
                count=count,
            ))
    if rng is not None:
        rng.shuffle(events)
    return events


def load_game_data(path: str, rng: Optional[random.Random] = None
                   ) -> Tuple[List[Character], Dict[Zone, List[Card]]]:
    return load_characters(path), load_decks(path, rng)
