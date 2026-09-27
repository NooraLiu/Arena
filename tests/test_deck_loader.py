import openpyxl
from arena.models import Zone, CardType
from arena.deck_loader import load_characters, load_decks, load_game_data


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    wc = wb.create_sheet("人物")
    wc.append(["ID", "角色名", "描述", "HP档(3/5/7)", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"])
    wc.append(["示例↓ C01", "示例·游侠", "x", 5, 3, "skill", 1, ""])   # example row, must be skipped
    wc.append(["C01", "游侠", "边境猎手", 5, 3, "可攻击相邻区域", 1, ""])
    wc.append(["C02", "刺客", "", 3, 4, "", 1, ""])

    ww = wb.create_sheet("武器")
    ww.append(["ID", "武器名", "描述", "攻击加成(+X)", "类型(装备/一次性)", "所在区域牌堆",
               "角色协同", "特殊效果", "张数", "备注"])
    ww.append(["示例↓ W01", "示例刀", "x", 2, "装备", "任意/外区", "", "", 3, ""])   # example
    ww.append(["W01", "猎刀", "短刀", 1, "装备", "外区", "刺客手中+2", "击杀抽食物", 2, ""])
    ww.append(["W02", "步枪", "强力", 4, "装备", "中心", "", "", 1, ""])

    ws = wb.create_sheet("物资(食物·护甲)")
    ws.append(["ID", "牌名", "描述", "类型(食物/护甲)", "效果", "数值", "所在区域牌堆",
               "特殊效果", "张数", "备注"])
    ws.append(["示例↓ S01", "蘑菇", "毒", "食物", "回血", 2, "任意", "", 4, ""])   # example
    ws.append(["S01", "面包", "", "食物", "回复HP", 2, "外区", "", 2, ""])
    ws.append(["S02", "盾", "", "护甲", "减伤", 2, "中心", "", 1, ""])
    wb.save(path)


def test_load_characters_skips_example_and_reads_stats(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_xlsx(p)
    chars = load_characters(str(p))
    assert [c.name for c in chars] == ["游侠", "刺客"]
    assert chars[0].hp_max == 5 and chars[0].base_attack == 3
    assert chars[1].hp_max == 3 and chars[1].base_attack == 4


def test_load_decks_places_cards_by_zone_with_counts(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_xlsx(p)
    decks = load_decks(str(p))
    # 猎刀: 外区 x2 -> 2 copies in EACH outer zone, none in center
    for z in [Zone.N, Zone.E, Zone.S, Zone.W]:
        knives = [c for c in decks[z] if c.name == "猎刀"]
        assert len(knives) == 2
        assert knives[0].type == CardType.WEAPON and knives[0].value == 1
    assert not any(c.name == "猎刀" for c in decks[Zone.CENTER])
    # 步枪: 中心 x1
    rifles = [c for c in decks[Zone.CENTER] if c.name == "步枪"]
    assert len(rifles) == 1 and rifles[0].value == 4
    # food / armor by type
    assert any(c.type == CardType.FOOD and c.name == "面包" for c in decks[Zone.N])
    assert any(c.type == CardType.ARMOR and c.name == "盾" for c in decks[Zone.CENTER])


def test_extra_columns_are_preserved(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_xlsx(p)
    decks = load_decks(str(p))
    knife = next(c for c in decks[Zone.N] if c.name == "猎刀")
    assert knife.description == "短刀"
    assert knife.synergy == "刺客手中+2"
    assert knife.effect == "击杀抽食物"


def test_zone_names_map_to_ring_positions():
    from arena.deck_loader import parse_zones
    # ring order 林区 -> 水区 -> 石区 -> 城区 -> (back to 林区)
    assert parse_zones("林区") == [Zone.N]
    assert parse_zones("水区") == [Zone.E]
    assert parse_zones("石区") == [Zone.S]
    assert parse_zones("城区") == [Zone.W]
    assert parse_zones("Central") == [Zone.CENTER]


def test_hp_column_named_plain_hp(tmp_path):
    p = tmp_path / "deck.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wc = wb.create_sheet("人物")
    wc.append(["ID", "角色名", "描述", "HP", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"])
    wc.append(["C06", "Cato", "", 15, 4, "", None, ""])
    wb.save(p)
    chars = load_characters(str(p))
    assert chars[0].name == "Cato" and chars[0].hp_max == 15 and chars[0].base_attack == 4


def test_load_game_data_returns_both(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_xlsx(p)
    chars, decks = load_game_data(str(p))
    assert len(chars) == 2 and Zone.CENTER in decks
