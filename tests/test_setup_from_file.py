import random
import openpyxl
from arena.models import Zone
from arena.setup import new_game
from arena.players.bots import RandomBot


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wc = wb.create_sheet("人物")
    wc.append(["ID", "角色名", "描述", "HP档(3/5/7)", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"])
    wc.append(["C01", "游侠", "", 5, 3, "", 1, ""])
    wc.append(["C02", "刺客", "", 3, 4, "", 1, ""])
    ww = wb.create_sheet("武器")
    ww.append(["ID", "武器名", "描述", "攻击加成(+X)", "类型(装备/一次性)", "所在区域牌堆",
               "角色协同", "特殊效果", "张数", "备注"])
    ww.append(["W02", "步枪", "", 4, "装备", "中心", "", "", 1, ""])
    ws = wb.create_sheet("物资(食物·护甲)")
    ws.append(["ID", "牌名", "描述", "类型(食物/护甲)", "效果", "数值", "所在区域牌堆",
               "特殊效果", "张数", "备注"])
    ws.append(["S01", "面包", "", "食物", "回血", 2, "外区", "", 2, ""])
    wb.save(path)


def test_new_game_uses_file_characters_and_decks(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_xlsx(p)
    rng = random.Random(0)
    eng = new_game(num_players=2, rng=rng,
                   players_factory=lambda: RandomBot(rng), data_path=str(p))
    assert {pl.character.name for pl in eng.state.players} == {"游侠", "刺客"}
    assert any(c.name == "步枪" for c in eng.state.decks[Zone.CENTER])


def test_new_game_without_path_still_uses_defaults(tmp_path):
    rng = random.Random(0)
    eng = new_game(num_players=3, rng=rng, players_factory=lambda: RandomBot(rng))
    assert len(eng.state.players) == 3


def test_empty_workbook_raises_clear_error(tmp_path):
    import pytest
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, header in [
        ("人物", ["ID", "角色名", "描述", "HP档(3/5/7)", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"]),
        ("武器", ["ID", "武器名", "描述", "攻击加成(+X)", "类型(装备/一次性)", "所在区域牌堆",
                  "角色协同", "特殊效果", "张数", "备注"]),
        ("物资(食物·护甲)", ["ID", "牌名", "描述", "类型(食物/护甲)", "效果", "数值",
                              "所在区域牌堆", "特殊效果", "张数", "备注"]),
    ]:
        wb.create_sheet(name).append(header)   # headers only, no data rows
    wb.save(p)
    rng = random.Random(0)
    with pytest.raises(ValueError, match="人物"):
        new_game(num_players=2, rng=rng,
                 players_factory=lambda: RandomBot(rng), data_path=str(p))
