import openpyxl
from arena.stats import run_many, summarize


def _make_filled_xlsx(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wc = wb.create_sheet("人物")
    wc.append(["ID", "角色名", "描述", "HP档(3/5/7)", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"])
    wc.append(["C01", "游侠", "", 5, 3, "", 1, ""])
    wc.append(["C02", "刺客", "", 3, 4, "", 1, ""])
    ww = wb.create_sheet("武器")
    ww.append(["ID", "武器名", "描述", "攻击加成(+X)", "类型(装备/一次性)", "所在区域牌堆",
               "角色协同", "特殊效果", "张数", "备注"])
    ww.append(["W01", "步枪", "", 4, "装备", "中心", "", "", 2, ""])
    ws = wb.create_sheet("物资(食物·护甲)")
    ws.append(["ID", "牌名", "描述", "类型(食物/护甲)", "效果", "数值", "所在区域牌堆",
               "特殊效果", "张数", "备注"])
    ws.append(["S01", "面包", "", "食物", "回血", 2, "外区", "", 2, ""])
    wb.save(path)


def test_run_many_with_deck_file(tmp_path):
    p = tmp_path / "deck.xlsx"
    _make_filled_xlsx(p)
    results = run_many(num_games=3, num_players=2, seed=0, data_path=str(p))
    assert len(results) == 3
    assert all(r["outcome"] in ("win", "draw", "capped") for r in results)


def test_run_many_and_summarize_shapes():
    results = run_many(num_games=20, num_players=6, seed=0)
    assert len(results) == 20
    s = summarize(results)
    assert s["avg_rounds"] >= 1
    assert set(s["outcome_counts"]).issubset({"win", "draw", "capped"})
    # win rates are fractions of all games and never exceed 1 in total
    assert sum(s["win_rate_by_seat"].values()) <= 1.0 + 1e-9
