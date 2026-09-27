"""Build small Arena deck workbooks for tests, in the per-zone layout."""
import openpyxl

CHAR_HEADER = ["ID", "角色名", "描述", "HP", "基础攻击(2/3/4)", "特殊技能", "张数", "备注"]
ZONE_HEADER = ["ID", "类型", "牌名", "描述", "数值", "角色协同", "特殊效果", "张数", "备注"]
ZONE_SHEETS = ["林区牌堆", "水区牌堆", "石区牌堆", "城区牌堆", "中心牌堆"]


def make_workbook(path, characters=None, zones=None):
    """characters: list of rows under CHAR_HEADER.
    zones: {sheet_name: [rows under ZONE_HEADER]}; every zone sheet is created (possibly empty)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wc = wb.create_sheet("人物")
    wc.append(CHAR_HEADER)
    for row in characters or []:
        wc.append(row)
    for sheet in ZONE_SHEETS:
        ws = wb.create_sheet(sheet)
        ws.append(ZONE_HEADER)
        for row in (zones or {}).get(sheet, []):
            ws.append(row)
    wb.save(path)
    return str(path)


DEFAULT_CHARS = [
    ["C01", "游侠", "", 13, 3, "", None, ""],
    ["C02", "刺客", "", 11, 4, "", None, ""],
]

DEFAULT_ZONES = {
    "中心牌堆": [["中01", "武器", "步枪", "", 4, "", "", 2, ""]],
    "林区牌堆": [["林01", "食物", "面包", "", 2, "", "", 2, ""]],
}
